#!/usr/bin/env python3
"""Generate M3 consent and data-flow documentation Word file for ABDM V3 HIP service."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "FlowExpalined-M3.docx"
ASSETS = Path(
    r"C:\Users\Lenovo\.cursor\projects\c-ShilpaWkspace-current-iplit-ABDM-repos-2-5-14-Upgraded-hip-service\assets"
)

NAVY = RGBColor(0x1B, 0x3A, 0x5F)
TEAL = RGBColor(0x0D, 0x73, 0x73)
DARK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_FILL = "1B3A5F"
ALT_ROW = "F3F7FA"


def set_run(run, size=11, bold=False, color=DARK, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, color=DARK, size=9, align="left"):
    cell.text = ""
    p = cell.paragraphs[0]
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)


def add_heading_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY if level <= 2 else TEAL
        run.font.name = "Calibri"
    return h


def add_para(doc, text, size=11, bold=False, color=DARK, space_after=8):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    run = p.add_run(text)
    set_run(run, size=11)
    p.paragraph_format.left_indent = Cm(1.25 + level * 0.6)
    p.paragraph_format.space_after = Pt(3)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run(run, size=9, font="Consolas", color=RGBColor(0x1A, 0x1A, 0x1A))
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.left_indent = Cm(0.4)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F4F6F8")
    shd.set(qn("w:val"), "clear")
    pPr.append(shd)
    return p


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, HEADER_FILL)
        set_cell_text(cell, h, bold=True, color=WHITE, size=9, align="center")
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            if r_idx % 2 == 1:
                shade_cell(cell, ALT_ROW)
            set_cell_text(cell, str(val), size=8.5)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def add_image(doc, filename, caption, width_inches=6.6):
    path = ASSETS / filename
    if not path.exists():
        add_para(doc, f"[Diagram not found: {filename}]", color=GRAY, size=10)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_inches))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    set_run(r, size=9, bold=False, color=GRAY)
    cap.paragraph_format.space_after = Pt(12)


def add_api_block(doc, title, rows):
    add_heading_styled(doc, title, 3)
    add_table(doc, ["Item", "Value"], rows, col_widths=[4.5, 12.5])


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("ABDM V3 — HIP Service API Documentation")
    set_run(r, size=26, bold=True, color=NAVY)
    t.paragraph_format.space_after = Pt(6)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run("Milestone 3 (M3): Consent and Health Information Exchange")
    set_run(r, size=16, bold=True, color=TEAL)

    for line in [
        "Service: hip-service (In.ProjectEKA.HipService) acting as HIP / data provider",
        "Companions: M1 docs/changesExpalined.docx , M2 docs/FlowExpalined-M2.docx",
        "Starting commit: 83ddacfe — “abdm V3”, plus data-transfer, hipId, doneAt, date-range, consent fixes",
        "References: ABDM sandbox V3 data-flow and consent paths",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        set_run(r, size=11, color=GRAY)
        p.paragraph_format.space_after = Pt(2)

    add_para(doc, "")

    add_heading_styled(doc, "1. What M3 covers", 1)
    add_para(
        doc,
        "Milestone 3 is health-information exchange. After M2 has linked care contexts to an ABHA address, a Health "
        "Information User (HIU) — typically a PHR app or another hospital — may request those records. The patient "
        "grants a consent artefact. ABDM notifies this HIP. Later ABDM asks this HIP to collect FHIR bundles from "
        "OpenMRS, encrypt them with the HIU’s ECDH key, push them to the HIU dataPushUrl, and notify Gateway of success or failure.",
    )
    add_table(
        doc,
        ["Step", "What happens", "HIP role"],
        [
            [
                "Consent notify",
                "Patient grants / revokes consent in PHR for this HIP’s care contexts",
                "Store or update artefact; acknowledge on-notify",
            ],
            [
                "Health-information request",
                "HIU asks ABDM for data against a consent id",
                "202, look up consent, enqueue RabbitMQ, on-request ACKNOWLEDGED or ERRORED",
            ],
            [
                "Collect + encrypt + push",
                "Background worker builds FHIR per hiType and care context",
                "ECDH/AES-GCM encrypt; POST to HIU dataPushUrl",
            ],
            [
                "Transfer notify",
                "ABDM and HIU need delivery status",
                "POST health-information/notify with hipId and per-context hiStatus",
            ],
        ],
        col_widths=[4.0, 7.0, 6.0],
    )
    add_para(
        doc,
        "This HIP is only the data provider. It does not initiate consent as an HIU. Legacy v0.5 routes "
        "POST /consent/notification and POST /health-information/request are marked [Obsolete]; V3 uses the /api/v3 paths below.",
    )

    add_heading_styled(doc, "2. Actors and prerequisites", 1)
    add_table(
        doc,
        ["Actor", "Role"],
        [
            ["Patient / PHR", "Grants or revokes consent for hiTypes and date range"],
            ["HIU", "Requests data; publishes ECDH public key + nonce; hosts dataPushUrl"],
            ["ABDM Gateway / HIE-CM", "Notifies HIP of consent; forwards HI request; receives on-request and notify"],
            ["hip-service", "Persists consent, collects FHIR, encrypts, pushes, notifies"],
            ["OpenMRS", "FHIR bundles via ws/rest/v1/hip/{hiType}/visit or /program"],
            ["RabbitMQ", "Async queue hip-data-request after HIP has acknowledged the HI request"],
            ["PostgreSQL", "ConsentArtefact, DataFlowRequest, HealthInformation (large-payload links)"],
        ],
        col_widths=[4.5, 12.5],
    )
    add_para(doc, "Prerequisites from M1/M2:", bold=True, space_after=4)
    add_bullet(doc, "Care contexts already linked to the ABHA address (M2). HIP resolves OpenMRS patientUuid from the linked account.")
    add_bullet(doc, "Care context referenceNumber format patientId:visitUuid so HFR id can be read from the visit location.")
    add_bullet(doc, "Gateway session token (same M1 pattern) on every HIP-to-Gateway and HIP-to-HIU call.")

    add_heading_styled(doc, "3. Flow A — Consent artefact notify", 1)
    add_para(
        doc,
        "When the patient grants consent for this HIP, Gateway POSTs the artefact. HIP stores it and must be able to "
        "look it up by consentId when the later health-information request arrives. Revoke / deny / expire update the same row.",
    )
    add_image(
        doc,
        "m3-consent-notify.png",
        "Figure 1. Consent notify: Gateway to HIP, persist artefact, on-notify acknowledgement.",
    )

    add_api_block(
        doc,
        "Consent notify (Gateway to HIP)",
        [
            ["Direction", "ABDM Gateway -> HIP callback"],
            ["HIP path", "POST /api/v3/consent/request/hip/notify"],
            ["HIP HTTP", "202 Accepted immediately; Hangfire StoreConsent"],
            ["HIP callback", "POST {Gateway.url}/api/hiecm/consent/v3/request/hip/on-notify"],
            ["Controller", "Consent.ConsentNotificationController"],
            ["Purpose", "Persist GRANTED artefact, or update REVOKED / DENIED / EXPIRED status"],
        ],
    )
    add_para(doc, "Inbound headers: CORRELATION-ID, REQUEST-ID (required), TIMESTAMP.", size=10)
    add_para(doc, "Gateway request body (ConsentArtefactRepresentation)", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "notification": {\n'
        '    "status": "GRANTED",\n'
        '    "consentId": "consent-uuid",\n'
        '    "signature": "<cms-signature>",\n'
        '    "grantAcknowledgement": true,\n'
        '    "consentDetail": {\n'
        '      "consentId": "consent-uuid",\n'
        '      "createdAt": "2026-08-22T06:00:00.000Z",\n'
        '      "patient": { "id": "first.last@sbx" },\n'
        '      "hip": { "id": "IN2710001820" },\n'
        '      "consentManager": { "id": "sbx" },\n'
        '      "hiTypes": ["OPConsultation", "Prescription"],\n'
        '      "careContexts": [\n'
        '        { "patientReference": "openMrsPatientUuid", "careContextReference": "openMrsPatientUuid:visitUuid" }\n'
        "      ],\n"
        '      "permission": {\n'
        '        "accessMode": "VIEW",\n'
        '        "dateRange": { "from": "2026-01-01T00:00:00.000Z", "to": "2026-08-22T23:59:59.000Z" },\n'
        '        "dataEraseAt": "2026-09-22T00:00:00.000Z"\n'
        "      }\n"
        "    }\n"
        "  }\n"
        "}",
    )
    add_para(doc, "HIP behaviour by status", bold=True, space_after=4)
    add_table(
        doc,
        ["notification.status", "HIP action", "on-notify"],
        [
            ["GRANTED", "Insert Consent (artefact id, detail, signature, status, consentId)", "Yes — acknowledgement.status OK"],
            ["REVOKED", "Update status; load artefact for cmSuffix", "Yes — acknowledgement.status OK"],
            ["DENIED / EXPIRED / REQUESTED", "Update status if row exists", "No callback in current code"],
        ],
        col_widths=[4.5, 7.5, 5.0],
    )
    add_para(doc, "HIP on-notify body (GatewayConsentRepresentation)", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "acknowledgement": {\n'
        '    "status": "OK",\n'
        '    "consentId": "consent-uuid"\n'
        "  },\n"
        '  "error": null,\n'
        '  "response": { "requestId": "<inbound REQUEST-ID>" }\n'
        "}",
    )
    add_para(
        doc,
        "ConsentStatus enum: REQUESTED, GRANTED, DENIED, REVOKED, EXPIRED. HIP does not verify the CMS signature in this service; "
        "it stores the signature with the artefact.",
        size=10,
        color=GRAY,
    )

    add_heading_styled(doc, "4. Flow B — Health information request and transfer", 1)
    add_para(
        doc,
        "After consent is GRANTED, the HIU asks ABDM for data. Gateway calls this HIP with the consent id, a date range, "
        "the HIU dataPushUrl, and ECDH keyMaterial. HIP acknowledges quickly, then a RabbitMQ worker collects OpenMRS FHIR, "
        "encrypts, pushes to the HIU, and notifies Gateway.",
    )
    add_image(
        doc,
        "m3-health-information-transfer.png",
        "Figure 2. Health-information request, collect FHIR, encrypt, push to HIU, notify Gateway.",
    )

    add_heading_styled(doc, "4.1 Receive request and on-request", 2)
    add_api_block(
        doc,
        "Health-information request (Gateway to HIP)",
        [
            ["Direction", "ABDM Gateway -> HIP callback"],
            ["HIP path", "POST /api/v3/hip/health-information/request"],
            ["HIP HTTP", "202 Accepted; Hangfire HealthInformationOf"],
            ["HIP callback", "POST {Gateway.url}/api/hiecm/data-flow/v3/health-information/hip/on-request"],
            ["Controller", "DataFlow.PatientDataFlowController"],
            ["Purpose", "Validate consent, persist request, enqueue collect/encrypt, acknowledge Gateway"],
        ],
    )
    add_para(doc, "Inbound headers: CORRELATION-ID, REQUEST-ID, TIMESTAMP, X-GatewayID (consent manager / gateway id).", size=10)
    add_para(doc, "Gateway request body (PatientHealthInformationRequest)", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "transactionId": "txn-uuid",\n'
        '  "hiRequest": {\n'
        '    "consent": { "id": "consent-uuid" },\n'
        '    "dateRange": {\n'
        '      "from": "2026-01-01T00:00:00.000Z",\n'
        '      "to": "2026-08-22T23:59:59.000Z"\n'
        "    },\n"
        '    "dataPushUrl": "https://hiu.example/api/v3/hiu/health-information/push",\n'
        '    "keyMaterial": {\n'
        '      "cryptoAlg": "ECDH",\n'
        '      "curve": "Curve25519",\n'
        '      "nonce": "<base64-32-bytes>",\n'
        '      "dhPublicKey": {\n'
        '        "expiry": "2026-08-23T00:00:00.000Z",\n'
        '        "parameters": "Curve25519/32byte random key",\n'
        '        "keyValue": "<base64-HIU-public-key>"\n'
        "      }\n"
        "    }\n"
        "  }\n"
        "}",
    )
    add_para(doc, "HIP processing before on-request:", bold=True, space_after=4)
    add_bullet(doc, "Load consent by hiRequest.consent.id. Missing artefact → error ContextArtefactIdNotFound.")
    add_bullet(
        doc,
        "Resolve OpenMRS patientUuid from linked account of consent.patient.id. If empty, wait 5 seconds and retry "
        "(covers race with M2 link-confirm).",
    )
    add_bullet(doc, "Save HealthInformationRequest keyed by transactionId. Duplicate/DB failure → internal error.")
    add_bullet(doc, "Reject if keyMaterial.dhPublicKey.expiry is already past (ErrorCode.ExpiredKeyPair).")
    add_bullet(doc, "Publish DataRequest to RabbitMQ exchange hip-data-request (topic).")
    add_para(doc, "HIP on-request body (success)", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "hiRequest": {\n'
        '    "transactionId": "txn-uuid",\n'
        '    "sessionStatus": "ACKNOWLEDGED"\n'
        "  },\n"
        '  "error": null,\n'
        '  "resp": { "requestId": "<inbound REQUEST-ID>" }\n'
        "}",
    )
    add_para(doc, "On failure sessionStatus is ERRORED and error.code/message are set. cmSuffix on the callback is Gateway.CmSuffix.", size=10)

    add_heading_styled(doc, "4.2 Collect FHIR from OpenMRS", 2)
    add_para(
        doc,
        "MessagingQueueListener consumes the queue and calls DataFlowMessageHandler. CollectHipService walks each "
        "granted care context and each consented hiType, then OpenMrsPatientData GETs Bahmni HIP REST.",
    )
    add_table(
        doc,
        ["hiType", "OpenMRS path root", "Query"],
        [
            ["OPConsultation", "ws/rest/v1/hip/opConsults/visit", "patientId, visitUuid, fromDate, toDate"],
            ["Prescription", "ws/rest/v1/hip/prescriptions/visit", "same"],
            ["DiagnosticReport", "ws/rest/v1/hip/diagnosticReports/visit", "same"],
            ["DischargeSummary", "ws/rest/v1/hip/dischargeSummary/visit", "same"],
            ["ImmunizationRecord", "ws/rest/v1/hip/immunizationRecord/visit", "same"],
            ["HealthDocumentRecord", "ws/rest/v1/hip/healthDocumentRecord/visit", "same"],
            ["WellnessRecord", "ws/rest/v1/hip/wellnessRecord/visit", "same"],
            ["Invoice", "ws/rest/v1/hip/invoice/visit", "same"],
        ],
        col_widths=[4.2, 7.0, 5.8],
    )
    add_para(
        doc,
        "Visit care contexts: visitUuid is the part after “:” in careContextReference. fromDate is expanded by minus one day "
        "and toDate by plus one day (V3 date-range fix). Program care contexts matching “(ID Number:n)” use "
        "ws/rest/v1/hip/{root}/program/ with programName and programEnrollmentId instead of visitUuid.",
        size=10,
        color=GRAY,
    )
    add_para(
        doc,
        "Each OpenMRS item’s bundle JSON is parsed as FHIR R4 Bundle and kept as CareBundle(careContextReference, bundleJson). "
        "Unknown hiType names are skipped.",
    )

    add_heading_styled(doc, "4.3 Encrypt (ECDH + AES-GCM)", 2)
    add_image(
        doc,
        "m3-encryption.png",
        "Figure 3. HIP derives a shared AES key with the HIU via ECDH, then AES-GCM encrypts each FHIR Bundle.",
    )
    add_para(
        doc,
        "DataEntryFactory generates an ephemeral EC key pair on the HIU’s curve (typically Curve25519) and a 32-byte HIP nonce. "
        "Encryptor derives a shared secret from the HIU dhPublicKey and HIP private key, XORs HIP nonce with HIU nonce, "
        "uses the first 20 bytes as salt and last 12 as IV, derives AES key, and encrypts UTF-8 FHIR JSON with AES-GCM (128-bit tag). "
        "Ciphertext is Base64. HIP returns its public key and nonce in the push payload so the HIU can decrypt.",
    )
    add_para(doc, "If Unicode byte size of ciphertext >= dataFlow.dataSizeLimitInMbs (default 5 MB):", bold=True, space_after=4)
    add_bullet(doc, "HIP stores the encrypted entry in HealthInformation with a random informationId and token.")
    add_bullet(doc, "Push payload contains a link instead of content: {hip.url}/health-information/{id}?token={token}.")
    add_bullet(doc, "HIU later GET that URL; HIP returns content if token matches and DataLinkTtlInMinutes (default 10) has not expired.")
    add_para(doc, "Checksum on entries is currently the literal string iplit-md5 when a bundle exists.", size=10, color=GRAY)

    add_heading_styled(doc, "4.4 Push encrypted data to HIU", 2)
    add_api_block(
        doc,
        "Data push (HIP to HIU)",
        [
            ["Direction", "HIP -> HIU dataPushUrl (not via HIE-CM body routing)"],
            ["HTTP", "POST {hiRequest.dataPushUrl}"],
            ["Auth", "Gateway session token (Authorization Bearer) plus REQUEST-ID, TIMESTAMP, X-CM-ID, CORRELATION-ID"],
            ["Facility header", "CreateHttpRequestWithContentType sets X-HIU-ID to the visit HFR id"],
            ["Content-Type", "application/json (payload); entries.media is application/fhir+json"],
            ["Class", "DataFlow.DataFlowClient.SendDataToHiu"],
        ],
    )
    add_para(
        doc,
        "HFR id is taken from the first granted care context’s visitUuid (same M2 location-attribute cache). "
        "If missing, HIP loads it from OpenMRS or uses the default Bahmni hip id. Transaction_Id header is the data transactionId.",
        size=10,
    )
    add_para(doc, "HIP push body (DataResponse)", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "pageNumber": 0,\n'
        '  "pageCount": 1,\n'
        '  "transactionId": "txn-uuid",\n'
        '  "entries": [\n'
        "    {\n"
        '      "content": "<base64-AES-GCM-ciphertext>",\n'
        '      "media": "application/fhir+json",\n'
        '      "checksum": "iplit-md5",\n'
        '      "careContextReference": "openMrsPatientUuid:visitUuid"\n'
        "    }\n"
        "  ],\n"
        '  "keyMaterial": {\n'
        '    "cryptoAlg": "ECDH",\n'
        '    "curve": "Curve25519",\n'
        '    "nonce": "<HIP-base64-nonce>",\n'
        '    "dhPublicKey": {\n'
        '      "expiry": "<copied from HIU key expiry>",\n'
        '      "parameters": "<copied from HIU>",\n'
        '      "keyValue": "<HIP-ephemeral-public-key-base64>"\n'
        "    }\n"
        "  }\n"
        "}",
    )

    add_heading_styled(doc, "4.5 Notify Gateway of transfer status", 2)
    add_api_block(
        doc,
        "Health-information notify (HIP to Gateway)",
        [
            ["Direction", "HIP -> Gateway"],
            ["Gateway path", "POST /api/hiecm/data-flow/v3/health-information/notify"],
            ["Class", "DataFlowNotificationClient.NotifyGateway"],
            ["Purpose", "Tell ABDM whether the session TRANSFERRED or FAILED, per care context"],
        ],
    )
    add_para(
        doc,
        "Always sent after the push attempt (success or catch). sessionStatus TRANSFERRED and hiStatus DELIVERED when the "
        "HTTP send to dataPushUrl succeeded and gateway authenticate succeeded; otherwise FAILED / ERRORED. "
        "notifier.type is HIP; notifier.id is Gateway.ClientId. hipId in statusNotification is the visit HFR id. "
        "doneAt is UTC DateTime.Today formatted yyyy-MM-ddTHH:mm:ss.fffZ.",
        size=10,
    )
    add_para(doc, "HIP notify body (GatewayDataNotificationRequest wrapper)", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "notification": {\n'
        '    "transactionId": "txn-uuid",\n'
        '    "consentId": "consent-uuid",\n'
        '    "doneAt": "2026-08-22T00:00:00.000Z",\n'
        '    "notifier": { "type": "HIP", "id": "<clientId>" },\n'
        '    "statusNotification": {\n'
        '      "sessionStatus": "TRANSFERRED",\n'
        '      "hipId": "IN2710001820",\n'
        '      "statusResponses": [\n'
        "        {\n"
        '          "careContextReference": "openMrsPatientUuid:visitUuid",\n'
        '          "hiStatus": "DELIVERED",\n'
        '          "description": "Successfully delivered health information"\n'
        "        }\n"
        "      ]\n"
        "    }\n"
        "  }\n"
        "}",
    )

    add_heading_styled(doc, "4.6 Optional: HIU fetches a large-payload link", 2)
    add_api_block(
        doc,
        "Health information by link (HIU to HIP)",
        [
            ["Direction", "HIU -> HIP"],
            ["HIP path", "GET /health-information/{informationId}?token={token}"],
            ["Controller", "DataFlow.DataFlowController.HealthInformation"],
            ["Purpose", "Return stored encrypted content when inline push was too large"],
        ],
    )
    add_table(
        doc,
        ["HTTP", "When"],
        [
            ["200 + HealthInformationResponse.content", "token matches and link not older than DataLinkTtlInMinutes"],
            ["403 InvalidToken / LinkExpired", "wrong token or TTL exceeded"],
            ["404 HealthInformationNotFound / ContextArtefactIdNotFound", "unknown id"],
            ["400 ExpiredKeyPair", "legacy mapping on this action result helper"],
        ],
        col_widths=[7.0, 10.0],
    )

    add_heading_styled(doc, "5. M3 API catalogue", 1)
    add_table(
        doc,
        ["#", "Path", "Direction", "Flow"],
        [
            ["A1", "POST /api/v3/consent/request/hip/notify", "GW -> HIP", "Consent"],
            ["A1b", "POST /api/hiecm/consent/v3/request/hip/on-notify", "HIP -> GW", "Consent"],
            ["B1", "POST /api/v3/hip/health-information/request", "GW -> HIP", "Data request"],
            ["B1b", "POST /api/hiecm/data-flow/v3/health-information/hip/on-request", "HIP -> GW", "Data request"],
            ["B2", "POST {dataPushUrl}", "HIP -> HIU", "Encrypted FHIR push"],
            ["B3", "POST /api/hiecm/data-flow/v3/health-information/notify", "HIP -> GW", "Transfer status"],
            ["B4", "GET /health-information/{id}?token=", "HIU -> HIP", "Large payload link"],
            ["L1", "POST /consent/notification [Obsolete]", "legacy", "Do not use for V3"],
            ["L2", "POST /health-information/request [Obsolete]", "legacy", "Do not use for V3"],
        ],
        col_widths=[1.4, 8.6, 3.2, 3.8],
    )

    add_heading_styled(doc, "6. Typical call order", 1)
    add_para(
        doc,
        "1. M2 has already linked the visit (careContextReference = patientId:visitUuid).\n"
        "2. Patient in PHR grants consent to an HIU for this HIP, selected hiTypes and date range.\n"
        "3. Gateway POST /api/v3/consent/request/hip/notify (GRANTED). HIP 202, saves artefact, on-notify OK.\n"
        "4. HIU requests data through ABDM. Gateway POST /api/v3/hip/health-information/request.\n"
        "5. HIP 202, on-request ACKNOWLEDGED, publishes RabbitMQ.\n"
        "6. Worker GETs OpenMRS FHIR per hiType, encrypts, POSTs DataResponse to dataPushUrl.\n"
        "7. HIP POST health-information/notify with TRANSFERRED / DELIVERED (or FAILED / ERRORED).\n"
        "8. If a later notify arrives with REVOKED, HIP updates artefact status and on-notify OK. Subsequent HI requests "
        "still find the row; they do not currently re-check GRANTED before collect — operationally ABDM should not request after revoke.",
    )

    add_heading_styled(doc, "7. Implementation notes from the V3 changes", 1)
    add_bullet(doc, "on-request and notify use /api/hiecm/data-flow/v3/... ; consent on-notify uses /api/hiecm/consent/v3/...")
    add_bullet(doc, "Data push and transfer notify include visit-wise HFR id (hipId), not only the default Bahmni id.")
    add_bullet(doc, "Data push uses CreateHttpRequestWithContentType, which sets header X-HIU-ID to that hipId.")
    add_bullet(doc, "doneAt on notify is UTC DateTime.Today in TIMESTAMP_FORMAT (doneAt fixes).")
    add_bullet(doc, "Visit FHIR fromDate/toDate are expanded by one day on each side.")
    add_bullet(doc, "If patientUuid is not yet in linked accounts, HIP waits 5 seconds once (link-confirm race).")
    add_bullet(doc, "RabbitMQ is required for actual transfer; acknowledging on-request without a running listener will not push data.")
    add_bullet(doc, "HIP never sends plaintext FHIR to the HIU; only AES-GCM ciphertext (or a short-lived link to ciphertext).")
    add_bullet(doc, "Obsolete v0.5 consent/notification and health-information/request remain compiled but are not the V3 contract.")

    add_heading_styled(doc, "8. Code map (M3)", 1)
    add_table(
        doc,
        ["Area", "Primary types"],
        [
            ["Consent notify", "Consent/ConsentNotificationController.cs, ConsentRepository, Common.Model.ConsentArtefact"],
            ["HI request", "DataFlow/DataFlowController.cs (PatientDataFlowController), DataFlow/DataFlow.cs"],
            ["Queue", "DataFlow/MessagingQueueListener.cs, DataFlowMessageHandler.cs"],
            ["Collect FHIR", "DataFlow/CollectHipService.cs, OpenMrsPatientData.cs"],
            ["Encrypt", "DataFlow/DataEntryFactory.cs, Encryptor/Encryptor.cs, EncryptorHelper.cs"],
            ["Push + notify", "DataFlow/DataFlowClient.cs, DataFlowNotificationClient.cs"],
            ["HFR on transfer", "Common/Model/BahmniConfiguration.cs (same cache as M2)"],
            ["Paths", "Common/Constants.cs PATH_CONSENTS_HIP, PATH_CONSENT_ON_NOTIFY, PATH_HEALTH_INFORMATION_*"],
        ],
        col_widths=[5.0, 12.0],
    )

    add_heading_styled(doc, "9. Document set", 1)
    add_bullet(doc, "M1 ABHA create/verify — docs/changesExpalined.docx")
    add_bullet(doc, "M2 link care contexts — docs/FlowExpalined-M2.docx")
    add_bullet(doc, "M3 consent + data flow — this file docs/FlowExpalined-M3.docx")

    add_heading_styled(doc, "10. External references", 1)
    add_bullet(doc, "ABDM Sandbox V3: https://sandbox.abdm.gov.in/sandbox/v3/new-documentation")
    add_bullet(doc, "HIP callback: POST /api/v3/consent/request/hip/notify")
    add_bullet(doc, "HIP callback: POST /api/v3/hip/health-information/request")
    add_bullet(doc, "HIP to Gateway: /api/hiecm/consent/v3/request/hip/on-notify")
    add_bullet(doc, "HIP to Gateway: /api/hiecm/data-flow/v3/health-information/hip/on-request")
    add_bullet(doc, "HIP to Gateway: /api/hiecm/data-flow/v3/health-information/notify")

    footer = doc.add_paragraph()
    r = footer.add_run("End of Milestone 3. M1, M2 and M3 together are the ABDM V3 HIP path in this service.")
    set_run(r, size=10, color=GRAY)
    footer.paragraph_format.space_before = Pt(18)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
