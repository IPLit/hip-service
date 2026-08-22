#!/usr/bin/env python3
"""Generate M1 API documentation Word file for ABDM V3 HIP service changes."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(__file__).resolve().parent / "changesExpalined.docx"
ASSETS = Path(
    r"C:\Users\Lenovo\.cursor\projects\c-ShilpaWkspace-current-iplit-ABDM-repos-2-5-14-Upgraded-hip-service\assets"
)

NAVY = RGBColor(0x1B, 0x3A, 0x5F)
TEAL = RGBColor(0x0D, 0x73, 0x73)
DARK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_FILL = "1B3A5F"
ALT_ROW = "F3F7FA"


def set_run(run, size=11, bold=False, color=DARK, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, color=DARK, size=9, align="left"):
    cell.text = ""
    p = cell.paragraphs[0]
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)


def add_heading_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY if level <= 2 else TEAL
        run.font.name = "Calibri"
    return h


def add_para(doc, text, size=11, bold=False, color=DARK, space_after=8):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    run = p.add_run(text)
    set_run(run, size=11)
    p.paragraph_format.left_indent = Cm(1.25 + level * 0.6)
    p.paragraph_format.space_after = Pt(3)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run(run, size=9, font="Consolas", color=RGBColor(0x1A, 0x1A, 0x1A))
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.left_indent = Cm(0.4)
    # light background via shading on paragraph
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F4F6F8")
    shd.set(qn("w:val"), "clear")
    pPr.append(shd)
    return p


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, HEADER_FILL)
        set_cell_text(cell, h, bold=True, color=WHITE, size=9, align="center")
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            if r_idx % 2 == 1:
                shade_cell(cell, ALT_ROW)
            set_cell_text(cell, str(val), size=8.5)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def add_image(doc, filename, caption, width_inches=6.6):
    path = ASSETS / filename
    if not path.exists():
        add_para(doc, f"[Diagram not found: {filename}]", color=GRAY, size=10)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_inches))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    set_run(r, size=9, bold=False, color=GRAY)
    cap.paragraph_format.space_after = Pt(12)


def add_api_block(doc, title, method, hip_path, abha_path, abha_method, purpose):
    add_heading_styled(doc, title, 3)
    add_table(
        doc,
        ["Item", "Value"],
        [
            ["HIP method & path", f"{method} {hip_path}"],
            ["Auth to HIP", "Bahmni session (reporting_session / JSESSIONID)"],
            ["Downstream ABHA API", f"{abha_method} {{AbhaNumberServiceUrl}}{abha_path}"],
            ["Purpose", purpose],
        ],
        col_widths=[4.5, 12.5],
    )


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    # Cover
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("ABDM V3 — HIP Service API Documentation")
    set_run(r, size=26, bold=True, color=NAVY)
    t.paragraph_format.space_after = Pt(6)

    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run("Milestone 1 (M1): ABHA Creation, Verification and Profile")
    set_run(r, size=16, bold=True, color=TEAL)

    meta = [
        "Service: hip-service (In.ProjectEKA.HipService)",
        "Starting commit: 83ddacfe4fa59b044b23309114d505e4b7b82c07 — “abdm V3”",
        "Scope of this section: M1 only. M2 (linking) and M3 (data flow) will be added after recap.",
        "References: ABDM Sandbox V3 documentation, User-initiated linking swagger, HIP-initiated linking.",
    ]
    for line in meta:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        set_run(r, size=11, color=GRAY)
        p.paragraph_format.space_after = Pt(2)

    add_para(doc, "")

    add_heading_styled(doc, "1. How to read this document", 1)
    add_para(
        doc,
        "This file documents the ABDM V3 work in hip-service from commit 83ddacfe (“abdm V3”) and the follow-on "
        "fixes on the trialsv3 line. M1, M2 and M3 are written in that order. This first delivery is M1 only, "
        "so the frontend and HIP contract can be recapped before linking (M2) and health-information exchange (M3).",
    )
    add_bullet(doc, "HIP path is what Bahmni / the hospital UI calls on hip-service.")
    add_bullet(doc, "ABHA path is what hip-service calls on ABDM after it obtains a gateway session token.")
    add_bullet(
        doc,
        "Sensitive values (Aadhaar, mobile, OTP) leave the hospital UI in plain text and are RSA-encrypted inside HIP before they are sent to ABDM.",
    )
    add_bullet(
        doc,
        "txnId and user tokens are stored in HIP memory against the Bahmni session. The UI usually does not need to round-trip txnId, except on the search-by-mobile login APIs.",
    )

    add_heading_styled(doc, "2. ABDM milestones in this service", 1)
    add_table(
        doc,
        ["Milestone", "ABDM capability", "HIP modules", "Status in this file"],
        [
            [
                "M1",
                "ABHA number & ABHA address: create, search, login, profile, card",
                "CreationController, VerificationController, EncryptionService, GatewayClient",
                "Documented now",
            ],
            [
                "M2",
                "User-initiated linking and HIP-initiated linking of care contexts; SMS notify; scan-and-share",
                "Discovery, Link, CareContext, UserAuth, SmsNotification, Patient share",
                "Next, after M1 recap",
            ],
            [
                "M3",
                "Consent notify and encrypted FHIR health-information transfer (HIP as data provider)",
                "Consent, DataFlow",
                "After M2",
            ],
        ],
        col_widths=[2.2, 6.5, 5.3, 3.0],
    )

    add_heading_styled(doc, "3. Common ABDM call pattern", 1)
    add_para(
        doc,
        "Every M1 ABHA call from HIP follows the same outer pattern. The hospital UI never talks to ABDM directly. "
        "HIP is the adapter: it authenticates the Bahmni user, fetches a short-lived ABDM gateway token, encrypts "
        "PII, calls ABHA, and keeps transaction state in memory for that Bahmni session.",
    )
    add_image(
        doc,
        "m1-common-abdm-call-pattern.png",
        "Figure 1. Common ABDM V3 call pattern used by all M1 APIs.",
    )

    add_heading_styled(doc, "3.1 Actors and base URLs", 2)
    add_table(
        doc,
        ["Actor", "Role", "Configured as"],
        [
            ["Bahmni UI / HIP client", "Calls HIP /v3/hip/* APIs with Bahmni session cookie", "OpenMRS session"],
            ["hip-service", "Encrypts PII, maps payloads, stores txnId/token per session", "This service"],
            [
                "ABDM Gateway",
                "Issues accessToken for subsequent ABHA and HIECM calls",
                "Gateway.url + /api/hiecm/gateway/v3/sessions",
            ],
            [
                "ABHA Number service",
                "Enrollment, login, profile, ABHA card (M1 core)",
                "Gateway.abhaNumberServiceUrl (sandbox: https://abhasbx.abdm.gov.in/abha/api)",
            ],
            [
                "ABHA Address / PHR service",
                "ABHA-address search, login, phr-card, legacy phr flows",
                "Gateway.abhaAddressServiceUrl (sandbox: https://phrsbx.abdm.gov.in/api)",
            ],
            ["UIDAI", "Aadhaar OTP and e-KYC. HIP never calls UIDAI directly.", "Reached only by ABHA"],
        ],
        col_widths=[4.0, 7.0, 6.0],
    )

    add_heading_styled(doc, "3.2 Step-by-step for one ABHA call", 2)
    add_para(doc, "1. Client authentication to HIP", bold=True, space_after=4)
    add_para(
        doc,
        "CreationController and VerificationController are decorated with [Authorize(AuthenticationSchemes = bahmni)]. "
        "CustomAuthenticationHandler accepts reporting_session or OpenMRS JSESSIONID, validates it against OpenMRS "
        "whoami, and stores session_id on HttpContext.Items. That session_id is the key for all in-memory ABHA state.",
    )
    add_para(doc, "2. Gateway session token", bold=True, space_after=4)
    add_para(
        doc,
        "GatewayClient.Authenticate POSTs clientId, clientSecret and grantType=client_credentials to "
        "{Gateway.url}/api/hiecm/gateway/v3/sessions with headers REQUEST-ID, TIMESTAMP, X-CM-ID and optional "
        "CORRELATION-ID. The accessToken is sent as Authorization on the ABHA call. HIP does this on every ABHA request.",
    )
    add_code(
        doc,
        "POST {Gateway.url}/api/hiecm/gateway/v3/sessions\n"
        "{\n"
        '  "clientId": "<clientId>",\n'
        '  "clientSecret": "<clientSecret>",\n'
        '  "grantType": "client_credentials"\n'
        "}",
    )
    add_para(doc, "3. Encryption of Aadhaar / mobile / OTP", bold=True, space_after=4)
    add_para(
        doc,
        "On startup, EncryptionService loads the ABHA RSA public key from GET {AbhaNumberServiceUrl}/v3/profile/public/certificate. "
        "HIP encrypts plain text with RSA OAEP SHA-1 (RSA/ECB/OAEPWithSHA-1AndMGF1Padding) and Base64-encodes the result. "
        "The hospital UI must send plaintext; HIP encrypts before the ABHA request.",
    )
    add_para(doc, "4. Downstream ABHA HTTP call", bold=True, space_after=4)
    add_para(
        doc,
        "HttpRequestHelper always adds REQUEST-ID (UUID), TIMESTAMP (yyyy-MM-dd'T'HH:mm:ss.fff'Z'), Authorization, "
        "and X-CM-ID. Optional headers used in M1: CORRELATION-ID, X-Token (user token after login/enrol), "
        "T-token (legacy mobile-login token), Transaction_Id (used by profile login verify).",
    )
    add_para(doc, "5. Session dictionaries (CreationMap)", bold=True, space_after=4)
    add_table(
        doc,
        ["Dictionary", "Key", "Value stored", "Used by"],
        [
            ["TxnDictionary", "Bahmni session_id", "Latest ABDM txnId", "Almost every OTP step"],
            [
                "HealthIdNumberTokenDictionary",
                "session_id",
                "User X-token after enrol / login",
                "getAbhaProfile, getAbhaCard, verifyAbhaAccount",
            ],
            [
                "HealthIdLoginScopeDictionary",
                "session_id",
                "Scope list used on requestOtp",
                "verifyOtp / abhaAddress verifyOtp",
            ],
            [
                "HealthIdTokenDictionary",
                "session_id",
                "PHR Bearer token",
                "Legacy LinkPhrController (v1 phr)",
            ],
            [
                "VerifiedMobileTokenDictionary",
                "session_id",
                "T-token from v2 mobile login",
                "Legacy GET_AUTHORIZED_TOKEN",
            ],
        ],
        col_widths=[5.0, 3.2, 5.0, 3.8],
    )
    add_para(
        doc,
        "These maps are ConcurrentDictionary instances (thread-safe). They are in-process memory, not the database. "
        "A HIP restart or a different HIP replica will lose txnId/token for that session.",
        size=10,
        color=GRAY,
    )

    add_heading_styled(doc, "3.3 Common request headers (HIP → ABDM)", 2)
    add_table(
        doc,
        ["Header", "Required", "Source in HIP"],
        [
            ["Authorization", "Yes", "Bearer accessToken from gateway sessions"],
            ["REQUEST-ID", "Yes", "New GUID per call"],
            ["TIMESTAMP", "Yes", "UTC, format yyyy-MM-dd'T'HH:mm:ss.fff'Z'"],
            ["X-CM-ID", "Yes", "Gateway.cmSuffix (sbx in sandbox)"],
            ["CORRELATION-ID", "When UI sent it", "Copied from incoming HIP header"],
            ["X-Token", "After user login/enrol", "Bearer {user token} from HealthIdNumberTokenDictionary"],
            ["Transaction_Id", "profileLoginVerify", "txnId from search/OTP step"],
        ],
        col_widths=[4.0, 4.0, 9.0],
    )

    add_heading_styled(doc, "4. M1 functionality overview", 1)
    add_para(
        doc,
        "Milestone 1 is ABHA identity. A patient at the hospital can create a new ABHA number with Aadhaar OTP, "
        "optionally verify a communication mobile that is not the Aadhaar-linked mobile, pick an ABHA address, "
        "and download the ABHA card. An existing patient can search and log in with ABHA number, Aadhaar, mobile, "
        "or ABHA address, then fetch profile and card. HIP does not persist Aadhaar numbers.",
    )
    add_para(doc, "Implemented M1 flows in this codebase:", bold=True, space_after=4)
    add_bullet(doc, "Flow A — Create ABHA number via Aadhaar OTP, optional mobile verify, ABHA address, ABHA card.")
    add_bullet(doc, "Flow B — Login / verify an existing ABHA (ABHA number, Aadhaar, or mobile identifier).")
    add_bullet(doc, "Flow C — Search ABHA by mobile, then OTP login for a selected index (V3 search-abha APIs).")
    add_bullet(doc, "Flow D — Search and login with ABHA address (PHR address service).")
    add_para(
        doc,
        "Not implemented as first-class V3 HIP APIs (ABDM supports them, this service does not expose them on /v3/hip): "
        "driving-licence enrolment, biometric enrol/login, demo-auth V3 body, email verification, delete/deactivate/re-KYC. "
        "A V1 demo-auth route still exists at /v1/hid/benefit/createHealthId/demo/auth.",
        size=10,
        color=GRAY,
    )

    add_heading_styled(doc, "5. Flow A — Create ABHA via Aadhaar OTP", 1)
    add_para(
        doc,
        "Official ABDM V3 sequence: request Aadhaar OTP → enrol by Aadhaar (creates 14-digit ABHA) → if communication "
        "mobile is different, request mobile OTP and auth/byAbdm → get address suggestions → create ABHA address → "
        "download ABHA card. Consent code abha-enrollment version 1.4 is added by HIP on enrol.",
    )
    add_image(
        doc,
        "m1-create-abha-aadhaar-otp.png",
        "Figure 2. M1 Flow A — Create ABHA number via Aadhaar OTP, then address and card.",
    )

    add_api_block(
        doc,
        "5.1 Generate Aadhaar OTP",
        "POST",
        "/v3/hip/generateAadhaarOtp",
        "/v3/enrollment/request/otp",
        "POST",
        "Send OTP to the mobile linked with Aadhaar. Stores txnId in TxnDictionary. Returns only the ABDM message to the UI (txnId is not returned).",
    )
    add_para(doc, "HIP request body", bold=True, space_after=4)
    add_code(doc, '{\n  "aadhaar": "123412341234"   // 12-digit Aadhaar, plaintext; HIP encrypts\n}')
    add_para(doc, "HIP → ABHA body (after encryption)", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "txnId": "",\n'
        '  "scope": ["abha-enrol"],\n'
        '  "loginHint": "aadhaar",\n'
        '  "loginId": "<RSA-encrypted Aadhaar>",\n'
        '  "otpSystem": "aadhaar"\n'
        "}",
    )
    add_para(doc, "HIP response (202 Accepted on success)", bold=True, space_after=4)
    add_code(doc, '{\n  "message": "OTP sent to Aadhaar registered mobile ending xx001"\n}')
    add_para(
        doc,
        "Resend OTP: call the same HIP API again. ABDM allows resend; the sandbox guidance is to enable resend at least twice after 60 seconds. "
        "HIP currently starts a new enrolment OTP (empty txnId) rather than passing the previous txnId for resend.",
        size=10,
        color=GRAY,
    )

    add_api_block(
        doc,
        "5.2 Verify Aadhaar OTP and create ABHA",
        "POST",
        "/v3/hip/verifyOtpAndCreateABHA",
        "/v3/enrollment/enrol/byAadhaar",
        "POST",
        "Verifies Aadhaar OTP, creates or returns the ABHA profile, stores user token for card/profile.",
    )
    add_para(doc, "HIP request body", bold=True, space_after=4)
    add_code(doc, '{\n  "otp": "123456",\n  "mobile": "98xxxxxxxx"   // communication mobile; sent unencrypted to ABHA in current code\n}')
    add_para(doc, "HIP → ABHA body", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "authData": {\n'
        '    "authMethods": ["otp"],\n'
        '    "otp": {\n'
        '      "txnId": "<from session>",\n'
        '      "otpValue": "<RSA-encrypted OTP>",\n'
        '      "mobile": "98xxxxxxxx",\n'
        '      "timestamp": "2026-08-22T06:25:00.000Z"\n'
        "    }\n"
        "  },\n"
        '  "consent": {\n'
        '    "code": "abha-enrollment",\n'
        '    "version": "1.4"\n'
        "  }\n"
        "}",
    )
    add_para(doc, "HIP response (200 OK)", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "message": "Account created successfully",\n'
        '  "isNew": true,\n'
        '  "abhaProfile": {\n'
        '    "firstName": "First",\n'
        '    "middleName": "",\n'
        '    "lastName": "Last",\n'
        '    "dob": "01-01-1990",\n'
        '    "gender": "F",\n'
        '    "mobile": "98xxxxxxxx",\n'
        '    "phrAddress": ["first.last@sbx"],\n'
        '    "address": "...",\n'
        '    "stateName": "...",\n'
        '    "districtName": "...",\n'
        '    "pinCode": "400001",\n'
        '    "abhaNumber": "91-xxxx-xxxx-xxxx",\n'
        '    "abhaStatus": "ACTIVE"\n'
        "  }\n"
        "}",
    )
    add_para(
        doc,
        "HIP stores enroll tokens.token in HealthIdNumberTokenDictionary and the new txnId in TxnDictionary. "
        "isNew=false means this Aadhaar already had an ABHA; the same profile is returned.",
        size=10,
        color=GRAY,
    )

    add_api_block(
        doc,
        "5.3 Generate mobile OTP (only if communication mobile ≠ Aadhaar mobile)",
        "POST",
        "/v3/hip/generateMobileOtp",
        "/v3/enrollment/request/otp",
        "POST",
        "OTP to the communication mobile. Uses existing txnId from session. Scopes abha-enrol + mobile-verify.",
    )
    add_para(doc, "HIP request body", bold=True, space_after=4)
    add_code(doc, '{\n  "mobile": "98xxxxxxxx"\n}')
    add_para(doc, "HIP → ABHA body", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "txnId": "<session txnId>",\n'
        '  "scope": ["abha-enrol", "mobile-verify"],\n'
        '  "loginHint": "mobile",\n'
        '  "loginId": "<RSA-encrypted mobile>",\n'
        '  "otpSystem": "abdm"\n'
        "}",
    )
    add_para(doc, "HIP response (202 Accepted)", bold=True, space_after=4)
    add_code(doc, '{\n  "message": "OTP sent to mobile ending xx89"\n}')

    add_api_block(
        doc,
        "5.4 Verify mobile OTP and link mobile",
        "POST",
        "/v3/hip/verifyMobileOtp",
        "/v3/enrollment/auth/byAbdm",
        "POST",
        "Links the verified communication mobile to the ABHA enrolment transaction.",
    )
    add_para(doc, "HIP request body", bold=True, space_after=4)
    add_code(doc, '{\n  "otp": "123456"\n}')
    add_para(doc, "HIP → ABHA body", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "scope": ["abha-enrol", "mobile-verify"],\n'
        '  "authData": {\n'
        '    "authMethods": ["otp"],\n'
        '    "otp": {\n'
        '      "txnId": "<session txnId>",\n'
        '      "otpValue": "<RSA-encrypted OTP>",\n'
        '      "timestamp": "2026-08-22T06:26:00.000Z"\n'
        "    }\n"
        "  }\n"
        "}",
    )
    add_para(doc, "HIP response (200 OK) — EnrollmentAuthByAbdmResponse", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "txnId": "...",\n'
        '  "authResult": "success",\n'
        '  "message": "Mobile verified successfully"\n'
        "}",
    )

    add_api_block(
        doc,
        "5.5 Get ABHA address suggestions",
        "GET",
        "/v3/hip/getAbhaAddressSuggestions",
        "/v3/enrollment/enrol/suggestion",
        "GET",
        "Returns suggested ABHA addresses for the current enrolment txnId (sent as Transaction_Id header).",
    )
    add_para(doc, "HIP request: no body. txnId is taken from session.", bold=True, space_after=4)
    add_para(doc, "HIP response (200 OK) — txnId is stripped; UI only receives the list", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "abhaAddressList": [\n'
        '    "first.last",\n'
        '    "firstlast90",\n'
        '    "first.l90"\n'
        "  ]\n"
        "}",
    )

    add_api_block(
        doc,
        "5.6 Create ABHA address",
        "POST",
        "/v3/hip/createAbhaAddress",
        "/v3/enrollment/enrol/abha-address",
        "POST",
        "Creates the preferred ABHA address for the enrolment transaction.",
    )
    add_para(doc, "HIP request body", bold=True, space_after=4)
    add_code(doc, '{\n  "abhaAddress": "first.last"\n}')
    add_para(doc, "HIP → ABHA body", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "txnId": "<session txnId>",\n'
        '  "abhaAddress": "first.last",\n'
        '  "preferred": 1\n'
        "}",
    )
    add_para(doc, "HIP response: 200 OK with empty body on success. Error body is forwarded from ABDM.", size=10)

    add_api_block(
        doc,
        "5.7 Download ABHA card",
        "GET",
        "/v3/hip/getAbhaCard",
        "/v3/profile/account/abha-card",
        "GET",
        "Returns PNG of the official ABHA card. Requires X-Token from enrol/login stored in session.",
    )
    add_para(
        doc,
        "HIP response: image/png stream. The UI should treat this as a file download. "
        "If the session has no token, this call fails with HTTP 500 in the current implementation.",
        size=10,
        color=GRAY,
    )

    add_heading_styled(doc, "6. Flow B — Verify / login existing ABHA", 1)
    add_para(
        doc,
        "Use this when the patient already has an ABHA. HIP maps identifierType + authMethod to ABDM scope, loginHint and otpSystem. "
        "After OTP success, HIP stores the user token. If several accounts are returned, the UI must call verifyAbhaAccount with the chosen ABHA number.",
    )
    add_image(
        doc,
        "m1-verify-login-abha.png",
        "Figure 3. M1 Flows B, C and D — login by ABHA/Aadhaar, search-by-mobile, and ABHA address.",
    )

    add_heading_styled(doc, "6.1 Identifier and auth mapping (requestOtp)", 2)
    add_table(
        doc,
        ["identifierType", "authMethod", "ABDM scope", "loginHint", "otpSystem"],
        [
            ["ABHA_NUMBER", "AADHAAR_OTP", "abha-login, aadhaar-verify", "abha-number", "aadhaar"],
            ["ABHA_NUMBER", "MOBILE_OTP", "abha-login, mobile-verify", "abha-number", "abdm"],
            ["AADHAAR_NUMBER", "(any / Aadhaar)", "abha-login, aadhaar-verify", "aadhaar", "aadhaar"],
            [
                "MOBILE_NUMBER",
                "(mapped)",
                "abha-login, mobile-verify, search-abha",
                "index",
                "abdm",
            ],
        ],
        col_widths=[3.4, 3.0, 5.2, 2.8, 2.6],
    )
    add_para(
        doc,
        "MOBILE_NUMBER on /verification/requestOtp uses loginHint=index. That path expects the identifier to be an ABHA list index, "
        "not a raw mobile number. For “find ABHA by mobile number”, use Flow C (searchByMobile) instead.",
        size=10,
        color=GRAY,
    )

    add_api_block(
        doc,
        "6.2 Request OTP",
        "POST",
        "/v3/hip/verification/requestOtp",
        "/v3/profile/login/request/otp",
        "POST",
        "Starts ABHA login. Stores txnId and scope against the Bahmni session.",
    )
    add_para(doc, "HIP request body", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "identifier": "91-xxxx-xxxx-xxxx",\n'
        '  "identifierType": "ABHA_NUMBER",\n'
        '  "authMethod": "AADHAAR_OTP"\n'
        "}",
    )
    add_para(doc, "Allowed identifierType: ABHA_NUMBER | MOBILE_NUMBER | AADHAAR_NUMBER. Allowed authMethod: MOBILE_OTP | AADHAAR_OTP.", size=10)
    add_para(doc, "HIP → ABHA body example (ABHA number + Aadhaar OTP)", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "scope": ["abha-login", "aadhaar-verify"],\n'
        '  "loginHint": "abha-number",\n'
        '  "loginId": "<RSA-encrypted ABHA number>",\n'
        '  "otpSystem": "aadhaar"\n'
        "}",
    )
    add_para(doc, "HIP response (202 Accepted)", bold=True, space_after=4)
    add_code(doc, '{\n  "message": "OTP sent to Aadhaar registered mobile ending xx001"\n}')
    add_para(doc, "Invalid identifierType or authMethod returns 400 with the exception message.", size=10)

    add_api_block(
        doc,
        "6.3 Verify OTP",
        "POST",
        "/v3/hip/verification/verifyOtp",
        "/v3/profile/login/verify",
        "POST",
        "Completes login. Stores user token. Returns auth result and linked accounts (token itself is not returned to UI).",
    )
    add_para(doc, "HIP request body", bold=True, space_after=4)
    add_code(doc, '{\n  "otp": "123456"\n}')
    add_para(doc, "HIP → ABHA body", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "scope": ["abha-login", "aadhaar-verify"],\n'
        '  "authData": {\n'
        '    "authMethods": ["otp"],\n'
        '    "otp": {\n'
        '      "txnId": "<session txnId>",\n'
        '      "otpValue": "<RSA-encrypted OTP>"\n'
        "    }\n"
        "  }\n"
        "}",
    )
    add_para(doc, "HIP response (200 OK) — VerificationVerifyOtpResponse", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "txnId": "...",\n'
        '  "authResult": "success",\n'
        '  "message": "OTP verified",\n'
        '  "accounts": [\n'
        "    {\n"
        '      "abhaNumber": "91-xxxx-xxxx-xxxx",\n'
        '      "preferredAbhaAddress": "first.last@sbx",\n'
        '      "name": "First Last",\n'
        '      "gender": "F",\n'
        '      "dob": "01-01-1990",\n'
        '      "status": "ACTIVE",\n'
        '      "kycVerified": true\n'
        "    }\n"
        "  ]\n"
        "}",
    )

    add_api_block(
        doc,
        "6.4 Verify / select ABHA account (when multiple accounts)",
        "POST",
        "/v3/hip/verification/verifyAbhaAccount",
        "/v3/profile/login/verify/user",
        "POST",
        "Selects one ABHA from the list returned by verify OTP. Replaces session user token.",
    )
    add_para(doc, "HIP request body", bold=True, space_after=4)
    add_code(doc, '{\n  "abhaNumber": "91-xxxx-xxxx-xxxx"\n}')
    add_para(
        doc,
        "HIP adds session txnId and sends X-Token from the previous verify step. Success is 200 OK with empty body. "
        "Subsequent getAbhaProfile / getAbhaCard use the new token.",
        size=10,
    )

    add_api_block(
        doc,
        "6.5 Get ABHA profile",
        "GET",
        "/v3/hip/verification/getAbhaProfile",
        "/v3/profile/account",
        "GET",
        "Full ABHA profile for the logged-in user. Requires session token from verify OTP or enrol.",
    )
    add_para(
        doc,
        "If the session has no token, HIP returns 401 with message: Session token not found. Please complete OTP verification first.",
        size=10,
    )
    add_para(doc, "HIP response (200 OK) — selected fields", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "abhaNumber": "91-xxxx-xxxx-xxxx",\n'
        '  "preferredAbhaAddress": "first.last@sbx",\n'
        '  "mobile": "98xxxxxxxx",\n'
        '  "firstName": "First",\n'
        '  "lastName": "Last",\n'
        '  "name": "First Last",\n'
        '  "yearOfBirth": "1990",\n'
        '  "monthOfBirth": "01",\n'
        '  "dayOfBirth": "01",\n'
        '  "gender": "F",\n'
        '  "address": "...",\n'
        '  "stateName": "...",\n'
        '  "districtName": "...",\n'
        '  "pincode": "400001",\n'
        '  "kycVerified": "true",\n'
        '  "authMethods": ["AADHAAR_OTP", "MOBILE_OTP"]\n'
        "}",
    )

    add_heading_styled(doc, "7. Flow C — Search ABHA by mobile (V3)", 1)
    add_para(
        doc,
        "This is the dedicated V3 “find ABHA using mobile” path added in the M1 search-by-mobile commits. "
        "It is the correct flow when the patient knows the mobile number but not the ABHA number. "
        "HIP encrypts mobile and loginId, stores txnId, and returns a simplified list (index, masked ABHA, name, gender).",
    )

    add_api_block(
        doc,
        "7.1 Search ABHA by mobile",
        "POST",
        "/v3/hip/verification/abha/searchByMobile",
        "/v3/profile/account/abha/search",
        "POST",
        "Returns masked ABHA numbers linked to the mobile. Stores txnId in session.",
    )
    add_para(doc, "HIP request body", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "scope": ["search-abha"],\n'
        '  "mobile": "98xxxxxxxx"   // plaintext; HIP encrypts before ABHA call\n'
        "}",
    )
    add_para(doc, "Validation: scope and mobile are required; otherwise 400 “scope and mobile are required.”", size=10)
    add_para(doc, "ABDM raw response (array) is mapped by HIP to:", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "txnId": "uuid-from-abdm",\n'
        '  "abhaList": [\n'
        "    {\n"
        '      "index": 1,\n'
        '      "abhaNumber": "xx-xxxx-xxxx-6514",\n'
        '      "name": "First Last",\n'
        '      "gender": "F"\n'
        "    }\n"
        "  ]\n"
        "}",
    )
    add_para(doc, "UI should display abhaList and let the patient pick an index for the next call.", size=10)

    add_api_block(
        doc,
        "7.2 Request OTP for selected ABHA index",
        "POST",
        "/v3/hip/verification/abha/profileLoginRequestOtp",
        "/v3/profile/login/request/otp",
        "POST",
        "Requests OTP for the ABHA chosen from search. loginId is encrypted. Returns txnId + message.",
    )
    add_para(doc, "HIP request body", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "scope": ["abha-login", "search-abha"],\n'
        '  "loginHint": "index",\n'
        '  "loginId": "1",\n'
        '  "otpSystem": "abdm",\n'
        '  "txnId": "<txnId from searchByMobile>"\n'
        "}",
    )
    add_para(
        doc,
        "Required fields: scope, loginHint, loginId, txnId. HIP encrypts loginId (the index). "
        "Unlike most M1 OTP APIs, this one returns txnId to the UI.",
        size=10,
    )
    add_para(doc, "HIP response (202 Accepted)", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "txnId": "uuid-from-abdm",\n'
        '  "message": "OTP sent to mobile ending xx89"\n'
        "}",
    )

    add_api_block(
        doc,
        "7.3 Profile login verify OTP",
        "POST",
        "/v3/hip/verification/abha/profileLoginVerify",
        "/v3/profile/login/verify",
        "POST",
        "Verifies OTP for the selected ABHA. Stores user token. Returns ABDM auth payload including token.",
    )
    add_para(doc, "HIP request body", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "scope": ["abha-login", "search-abha"],\n'
        '  "authData": {\n'
        '    "authMethods": ["otp"],\n'
        '    "otp": {\n'
        '      "txnId": "<txnId from previous step>",\n'
        '      "otpValue": "123456"\n'
        "    }\n"
        "  }\n"
        "}",
    )
    add_para(doc, "HIP encrypts otpValue and also sends Transaction_Id header = txnId.", size=10)
    add_para(doc, "HIP response (202 Accepted) — ABHALoginVerifyOTPResponse when authResult is not failed", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "txnId": "...",\n'
        '  "authResult": "success",\n'
        '  "message": "OTP verified successfully",\n'
        '  "token": "<user-token>",\n'
        '  "expiresIn": 1800,\n'
        '  "refreshToken": "...",\n'
        '  "refreshExpiresIn": 1296000,\n'
        '  "accounts": [ { "abhaNumber": "91-xxxx-xxxx-xxxx", "preferredAbhaAddress": "...", "name": "..." } ]\n'
        "}",
    )
    add_para(
        doc,
        "After success, GET /v3/hip/verification/getAbhaProfile and GET /v3/hip/getAbhaCard work with the stored session token. "
        "If authResult is failed, HIP forwards the ABDM status and body.",
        size=10,
        color=GRAY,
    )

    add_heading_styled(doc, "8. Flow D — ABHA address (PHR) login", 1)
    add_para(
        doc,
        "These APIs talk to AbhaAddressServiceUrl (PHR), not the ABHA number enrolment APIs. "
        "Use them when the patient has an ABHA address such as name@sbx.",
    )

    add_api_block(
        doc,
        "8.1 Search ABHA address",
        "POST",
        "/v3/hip/verification/abhaAddress/search",
        "/login/abha/search",
        "POST",
        "Looks up an ABHA address and returns available auth methods.",
    )
    add_para(doc, "HIP request body", bold=True, space_after=4)
    add_code(doc, '{\n  "abhaAddress": "first.last@sbx"\n}')
    add_para(doc, "HIP response (202 Accepted) — SearchAbhaAddressResponse", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "healthIdNumber": "91-xxxx-xxxx-xxxx",\n'
        '  "abhaAddress": "first.last@sbx",\n'
        '  "authMethods": ["AADHAAR_OTP", "MOBILE_OTP"],\n'
        '  "blockedAuthMethods": [],\n'
        '  "status": "ACTIVE",\n'
        '  "fullName": "First Last",\n'
        '  "mobile": "xxxxxx7890"\n'
        "}",
    )

    add_api_block(
        doc,
        "8.2 Request OTP for ABHA address",
        "POST",
        "/v3/hip/verification/abhaAddress/requestOtp",
        "/login/abha/request/otp",
        "POST",
        "Sends OTP for ABHA-address login. Scope is abha-address-login + aadhaar-verify or mobile-verify.",
    )
    add_para(doc, "HIP request body", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "abhaAddress": "first.last@sbx",\n'
        '  "authMethod": "MOBILE_OTP"\n'
        "}",
    )
    add_para(doc, "HIP response (202 Accepted): { \"message\": \"...\" }. txnId stored in session.", size=10)

    add_api_block(
        doc,
        "8.3 Verify OTP for ABHA address",
        "POST",
        "/v3/hip/verification/abhaAddress/verifyOtp",
        "/login/abha/verify",
        "POST",
        "Verifies OTP and stores PHR user token. Returns users list without the raw token.",
    )
    add_para(doc, "HIP request body", bold=True, space_after=4)
    add_code(doc, '{\n  "otp": "123456"\n}')
    add_para(doc, "HIP response (200 OK)", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "authResult": "success",\n'
        '  "message": "OTP verified",\n'
        '  "users": [\n'
        "    {\n"
        '      "abhaAddress": "first.last@sbx",\n'
        '      "fullName": "First Last",\n'
        '      "abhaNumber": "91-xxxx-xxxx-xxxx",\n'
        '      "status": "ACTIVE",\n'
        '      "kycStatus": "VERIFIED"\n'
        "    }\n"
        "  ]\n"
        "}",
    )

    add_api_block(
        doc,
        "8.4 Get ABHA address profile",
        "GET",
        "/v3/hip/verification/abhaAddress/getProfile",
        "/login/profile/abha-profile",
        "GET",
        "Profile for the logged-in ABHA address. Uses X-Token from verify step.",
    )
    add_para(doc, "HIP response (200 OK) includes abhaAddress, abhaNumber, demographics, kycStatus, mobileVerified.", size=10)

    add_api_block(
        doc,
        "8.5 Get ABHA address / PHR card",
        "GET",
        "/v3/hip/verification/abhaAddress/getCard",
        "/login/profile/abha/phr-card",
        "GET",
        "PNG PHR card for the logged-in ABHA address.",
    )
    add_para(doc, "HIP response: image/png stream.", size=10)

    add_heading_styled(doc, "9. M1 API catalogue (HIP facing)", 1)
    add_table(
        doc,
        ["#", "HIP API", "ABDM API", "Flow"],
        [
            ["A1", "POST /v3/hip/generateAadhaarOtp", "POST /v3/enrollment/request/otp", "Create"],
            ["A2", "POST /v3/hip/verifyOtpAndCreateABHA", "POST /v3/enrollment/enrol/byAadhaar", "Create"],
            ["A3", "POST /v3/hip/generateMobileOtp", "POST /v3/enrollment/request/otp", "Create (optional mobile)"],
            ["A4", "POST /v3/hip/verifyMobileOtp", "POST /v3/enrollment/auth/byAbdm", "Create (optional mobile)"],
            ["A5", "GET /v3/hip/getAbhaAddressSuggestions", "GET /v3/enrollment/enrol/suggestion", "Create"],
            ["A6", "POST /v3/hip/createAbhaAddress", "POST /v3/enrollment/enrol/abha-address", "Create"],
            ["A7", "GET /v3/hip/getAbhaCard", "GET /v3/profile/account/abha-card", "Create / after login"],
            ["B1", "POST /v3/hip/verification/requestOtp", "POST /v3/profile/login/request/otp", "Login"],
            ["B2", "POST /v3/hip/verification/verifyOtp", "POST /v3/profile/login/verify", "Login"],
            ["B3", "POST /v3/hip/verification/verifyAbhaAccount", "POST /v3/profile/login/verify/user", "Login (multi-account)"],
            ["B4", "GET /v3/hip/verification/getAbhaProfile", "GET /v3/profile/account", "Login"],
            ["C1", "POST /v3/hip/verification/abha/searchByMobile", "POST /v3/profile/account/abha/search", "Search by mobile"],
            ["C2", "POST /v3/hip/verification/abha/profileLoginRequestOtp", "POST /v3/profile/login/request/otp", "Search by mobile"],
            ["C3", "POST /v3/hip/verification/abha/profileLoginVerify", "POST /v3/profile/login/verify", "Search by mobile"],
            ["D1", "POST /v3/hip/verification/abhaAddress/search", "POST /login/abha/search", "ABHA address"],
            ["D2", "POST /v3/hip/verification/abhaAddress/requestOtp", "POST /login/abha/request/otp", "ABHA address"],
            ["D3", "POST /v3/hip/verification/abhaAddress/verifyOtp", "POST /login/abha/verify", "ABHA address"],
            ["D4", "GET /v3/hip/verification/abhaAddress/getProfile", "GET /login/profile/abha-profile", "ABHA address"],
            ["D5", "GET /v3/hip/verification/abhaAddress/getCard", "GET /login/profile/abha/phr-card", "ABHA address"],
        ],
        col_widths=[1.2, 6.6, 5.8, 3.4],
    )

    add_heading_styled(doc, "10. Typical UI call order", 1)
    add_heading_styled(doc, "10.1 New patient — create ABHA", 2)
    add_para(
        doc,
        "1. Collect consent (ABDM enrolment consent, version 1.4 is sent by HIP on enrol).\n"
        "2. POST /v3/hip/generateAadhaarOtp with aadhaar.\n"
        "3. Patient enters Aadhaar OTP and communication mobile → POST /v3/hip/verifyOtpAndCreateABHA.\n"
        "4. Show ABHA number from abhaProfile. If communication mobile must be verified separately → generateMobileOtp then verifyMobileOtp.\n"
        "5. GET /v3/hip/getAbhaAddressSuggestions → patient picks or types an address → POST /v3/hip/createAbhaAddress.\n"
        "6. GET /v3/hip/getAbhaCard to display / download PNG.\n"
        "7. Continue hospital registration using abhaNumber and preferred ABHA address (M2 linking is a later step).",
    )

    add_heading_styled(doc, "10.2 Existing patient — knows ABHA number", 2)
    add_para(
        doc,
        "1. POST /v3/hip/verification/requestOtp with identifierType=ABHA_NUMBER and AADHAAR_OTP or MOBILE_OTP.\n"
        "2. POST /v3/hip/verification/verifyOtp.\n"
        "3. If accounts.length > 1, POST /v3/hip/verification/verifyAbhaAccount with the chosen abhaNumber.\n"
        "4. GET /v3/hip/verification/getAbhaProfile and optionally GET /v3/hip/getAbhaCard.",
    )

    add_heading_styled(doc, "10.3 Existing patient — knows only mobile", 2)
    add_para(
        doc,
        "1. POST /v3/hip/verification/abha/searchByMobile with scope [\"search-abha\"] and mobile.\n"
        "2. Patient selects an index from abhaList.\n"
        "3. POST /v3/hip/verification/abha/profileLoginRequestOtp with loginHint=index, loginId=<index>, txnId from step 1.\n"
        "4. POST /v3/hip/verification/abha/profileLoginVerify with the OTP and txnId from step 3.\n"
        "5. GET /v3/hip/verification/getAbhaProfile.",
    )

    add_heading_styled(doc, "10.4 Existing patient — knows ABHA address", 2)
    add_para(
        doc,
        "1. POST /v3/hip/verification/abhaAddress/search.\n"
        "2. POST /v3/hip/verification/abhaAddress/requestOtp with the authMethod returned by search.\n"
        "3. POST /v3/hip/verification/abhaAddress/verifyOtp.\n"
        "4. GET /v3/hip/verification/abhaAddress/getProfile and optionally getCard.",
    )

    add_heading_styled(doc, "11. Error handling and operational notes", 1)
    add_bullet(doc, "ABDM non-success status codes and response bodies are forwarded as-is (HIP uses StatusCode((int)response.StatusCode, responseContent)).")
    add_bullet(doc, "Unhandled exceptions in HIP return HTTP 500 with no body.")
    add_bullet(doc, "Always send CORRELATION-ID from the UI; HIP copies it to gateway and ABHA calls for tracing.")
    add_bullet(doc, "Do not store Aadhaar in HMIS. HIP also does not persist Aadhaar; only encrypted transit to ABDM.")
    add_bullet(doc, "Session maps are in-memory. Keep create/login steps on the same HIP instance and the same Bahmni session cookie.")
    add_bullet(doc, "Public certificate is loaded once at HIP startup. ABDM rotates it periodically (order of months); restart HIP after rotation if encryption starts failing.")
    add_bullet(doc, "Legacy V1/V2 routes still exist (auth/init, mobile generateOtp, phr/login/mobileEmail/*, demo auth). Prefer the /v3/hip APIs above for new UI work.")

    add_heading_styled(doc, "12. Code map (M1)", 1)
    add_table(
        doc,
        ["Area", "Primary types"],
        [
            ["Create APIs", "Creation/CreationController.cs"],
            ["Verify / search / address APIs", "Verification/VerificationController.cs"],
            ["Identifier → ABDM mapping", "Verification/VerificationRequestMapper.cs"],
            ["Session maps", "Creation/CreationMap.cs"],
            ["RSA encrypt", "Common/EncryptionService.cs"],
            ["Gateway token + ABHA HTTP", "Gateway/GatewayClient.cs, Common/HttpRequestHelper.cs"],
            ["Paths", "Common/Constants.cs (APP_PATH_* and ABHA_*)"],
            ["Auth to HIP", "Common/Model/CustomAuthenticationHandler.cs"],
        ],
        col_widths=[5.5, 11.5],
    )

    add_heading_styled(doc, "13. What comes next (not in this section)", 1)
    add_para(
        doc,
        "M2 will document user-initiated linking (discover → link init → link confirm), HIP-initiated linking "
        "(generate-token, carecontext, context notify, SMS notify2), scan-and-share (patient/share), and HFR id per care context. "
        "M3 will document consent/hip/notify and health-information request / encrypted FHIR push.",
    )
    add_para(
        doc,
        "Please recap M1 (API list, call order, and whether the UI should keep using session-hidden txnId versus the search-by-mobile APIs that return txnId). After that, M2 will be appended to this same file.",
        bold=True,
    )

    add_heading_styled(doc, "14. External references", 1)
    add_bullet(doc, "ABDM Sandbox V3 documentation: https://sandbox.abdm.gov.in/sandbox/v3/new-documentation")
    add_bullet(doc, "User-initiated linking swagger (M2): sandbox v3 swagger with integration_label abdm_user_initiated_linking_phr")
    add_bullet(doc, "HIP-initiated linking (M2): https://sandbox.abdm.gov.in/sandbox/v3/new-documentation?doc=HIPInitiatedlinking")
    add_bullet(doc, "ABHA number sandbox base: https://abhasbx.abdm.gov.in/abha/api")
    add_bullet(doc, "Gateway sessions: POST {Gateway.url}/api/hiecm/gateway/v3/sessions")

    footer = doc.add_paragraph()
    r = footer.add_run("End of Milestone 1. M2 and M3 sections will follow in this document after recap.")
    set_run(r, size=10, color=GRAY)
    footer.paragraph_format.space_before = Pt(18)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
