#!/usr/bin/env python3
"""Generate M1 change documentation: only iplit V3 deltas vs community abdm-v3."""

from pathlib import Path

from docx.shared import Pt

from _doc_style import (
    GRAY,
    add_api_block,
    add_bullet,
    add_code,
    add_cover,
    add_heading_styled,
    add_para,
    add_table,
    new_doc,
    set_run,
)

OUT = Path(__file__).resolve().parent.parent / "V3ChangesExpalined-M1.docx"


def build():
    doc = new_doc()
    add_cover(
        doc,
        "ABDM V3 — Changes on top of Bahmni community hip-service",
        "Milestone 1 (M1): ABHA creation, verification and search-by-mobile",
        [
            "Baseline (current folder): BahmniCommunity/hip-service branch abdm-v3",
            "Changed code: ABDM-repos-2.5.14-Upgraded/hip-service branch trialsv3-hfrid",
            "Starting commit in upgraded repo: 83ddacfe — “abdm V3” (author shilpa-iplit)",
            "This file documents only M1 deltas. Unchanged community M1 APIs are listed as “already present”.",
        ],
    )

    add_heading_styled(doc, "1. Scope of this document", 1)
    add_para(
        doc,
        "This is not a full M1 API manual. Community abdm-v3 already implements ABHA create, login by "
        "ABHA number/Aadhaar, ABHA-address login, profile and card. This file records only the work "
        "added in the upgraded hip-service after commit 83ddacfe, compared with the current community folder.",
    )
    add_bullet(doc, "New HIP APIs are documented end-to-end (request, ABDM mapping, response).")
    add_bullet(doc, "Existing HIP APIs appear only where behaviour, payload or error handling changed.")
    add_bullet(doc, "M2 (linking / HFR) and M3 (consent / data flow) are in V3ChangesExpalined-M2.docx and M3.")
    add_para(
        doc,
        "ABDM sandbox V3 documentation and swagger require a logged-in sandbox session. Paths below "
        "match the ABHA Number V3 APIs used by this service (abhasbx.abdm.gov.in/abha/api).",
        size=10,
        color=GRAY,
    )

    add_heading_styled(doc, "2. What community already had (not re-documented)", 1)
    add_table(
        doc,
        ["HIP API", "ABDM API", "Status in your branch"],
        [
            ["POST /v3/hip/generateAadhaarOtp", "POST /v3/enrollment/request/otp", "Unchanged contract; session map is now thread-safe"],
            ["POST /v3/hip/verifyOtpAndCreateABHA", "POST /v3/enrollment/enrol/byAadhaar", "Stores enrol txnId as well as token"],
            ["POST /v3/hip/generateMobileOtp", "POST /v3/enrollment/request/otp", "Unchanged contract"],
            ["POST /v3/hip/verifyMobileOtp", "POST /v3/enrollment/auth/byAbdm", "Unchanged contract"],
            ["GET /v3/hip/getAbhaAddressSuggestions", "GET /v3/enrollment/enrol/suggestion", "Updates session txnId from ABDM"],
            ["POST /v3/hip/createAbhaAddress", "POST /v3/enrollment/enrol/abha-address", "Safer txnId lookup"],
            ["GET /v3/hip/getAbhaCard", "GET /v3/profile/account/abha-card", "Safer token lookup"],
            ["POST /v3/hip/verification/requestOtp", "POST /v3/profile/login/request/otp", "MOBILE_NUMBER mapping changed (see §5)"],
            ["POST /v3/hip/verification/verifyOtp", "POST /v3/profile/login/verify", "JSON property names camelCase"],
            ["POST /v3/hip/verification/verifyAbhaAccount", "POST /v3/profile/login/verify/user", "Safer token lookup"],
            ["GET /v3/hip/verification/getAbhaProfile", "GET /v3/profile/account", "401 if no session token; profile path already V3"],
            ["ABHA address search / OTP / profile / card", "PHR /login/abha/*", "Thread-safe maps only"],
        ],
        col_widths=[5.4, 5.8, 5.8],
    )

    add_heading_styled(doc, "3. M1 change summary", 1)
    add_table(
        doc,
        ["Change", "Kind", "Why"],
        [
            [
                "Search ABHA by mobile (3 new HIP APIs)",
                "New",
                "ABDM V3 search uses POST /v3/profile/account/abha/search then login by list index. Community had no HIP wrapper.",
            ],
            [
                "ABHAScope.SEARCH_ABHA and ABHALoginHint.ABHA_INDEX",
                "New",
                "Required by ABDM search-abha + loginHint=index.",
            ],
            [
                "requestOtp with identifierType=MOBILE_NUMBER",
                "Behaviour change",
                "Now sends loginHint=index and scope search-abha. Do not send a raw mobile here; use searchByMobile.",
            ],
            [
                "CreationMap / Verification session dictionaries",
                "Hardening",
                "Dictionary → ConcurrentDictionary; AddOrUpdate / TryGetValue to avoid races and KeyNotFoundException.",
            ],
            [
                "Enrol-by-Aadhaar stores txnId",
                "Fix",
                "Address suggestion and create-address need the enrol txnId, not only the Aadhaar-OTP txnId.",
            ],
            [
                "getABHAProfile uses /v3/profile/account",
                "Fix",
                "Stopped calling legacy /v2/account/profile (ABHA_PATIENT_PROFILE).",
            ],
            [
                "getAbhaProfile returns 401 when token missing",
                "Fix",
                "Community indexed the map and could 500. UI now gets a clear message.",
            ],
            [
                "ABHALoginVerifyOTPRequest camelCase JSON",
                "Fix",
                "ABDM V3 expects authData / otp / txnId / otpValue, not PascalCase.",
            ],
            [
                "ABHA card / PHR card token lookup",
                "Fix",
                "TryGetValue instead of HealthIdNumberTokenDictionary[sessionId] (KeyNotFound → 500).",
            ],
        ],
        col_widths=[5.2, 3.0, 8.8],
    )

    add_heading_styled(doc, "4. New flow — Search ABHA by mobile", 1)
    add_para(
        doc,
        "Use this when the patient knows the mobile number but not the 14-digit ABHA number. "
        "ABDM returns a masked list. The UI picks an index, HIP requests OTP for that index, then verifies OTP. "
        "This matches ABDM V3 POST /v3/profile/account/abha/search followed by /v3/profile/login/request/otp "
        "with loginHint=index and scope [abha-login, search-abha].",
    )
    add_para(doc, "Call order", bold=True, space_after=4)
    add_para(
        doc,
        "1. POST /v3/hip/verification/abha/searchByMobile  →  patient picks index from abhaList\n"
        "2. POST /v3/hip/verification/abha/profileLoginRequestOtp  (loginId = index, txnId from step 1)\n"
        "3. POST /v3/hip/verification/abha/profileLoginVerify  (OTP + txnId from step 2)\n"
        "4. GET /v3/hip/verification/getAbhaProfile and optionally GET /v3/hip/getAbhaCard",
    )
    add_para(
        doc,
        "HIP encrypts mobile, loginId (index) and otpValue with the ABHA RSA certificate before calling ABDM. "
        "The UI must send plaintext. Auth to HIP is the same Bahmni session cookie as other /v3/hip APIs.",
        size=10,
        color=GRAY,
    )

    add_api_block(
        doc,
        "4.1 POST /v3/hip/verification/abha/searchByMobile  (NEW)",
        [
            ["HIP method & path", "POST /v3/hip/verification/abha/searchByMobile"],
            ["Auth to HIP", "Bahmni session (reporting_session / JSESSIONID)"],
            ["Downstream ABHA", "POST {AbhaNumberServiceUrl}/v3/profile/account/abha/search"],
            ["Controller", "VerificationController.SearchAbhaByMobile"],
            ["Constant", "APP_PATH_ABHA_SEARCH_BY_MOBILE / ABHA_SEARCH_BY_MOBILE"],
            ["Purpose", "Return masked ABHA numbers linked to a mobile. Store txnId in session."],
        ],
    )
    add_para(doc, "HIP request body (plaintext mobile; HIP encrypts)", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "scope": ["search-abha"],\n'
        '  "mobile": "98xxxxxxxx"\n'
        "}",
    )
    add_para(doc, "Validation: scope and mobile are required; otherwise HTTP 400 “scope and mobile are required.”", size=10)
    add_para(doc, "HIP → ABHA: same JSON after RSA-encrypting mobile.", size=10)
    add_para(
        doc,
        "ABDM returns an array. HIP takes the first element, reads txnId and ABHA[], and maps to a simple list "
        "(index, ABHANumber, name, gender). txnId is stored in TxnDictionary[session_id].",
        size=10,
    )
    add_para(doc, "HIP response (200 OK)", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "txnId": "uuid-from-abdm",\n'
        '  "abhaList": [\n'
        "    {\n"
        '      "index": 1,\n'
        '      "abhaNumber": "xx-xxxx-xxxx-6514",\n'
        '      "name": "First Last",\n'
        '      "gender": "F"\n'
        "    }\n"
        "  ]\n"
        "}",
    )
    add_para(doc, "Empty ABDM array → 200 with empty SearchAbhaByMobileResponse. Non-success ABDM status is forwarded.", size=10)

    add_api_block(
        doc,
        "4.2 POST /v3/hip/verification/abha/profileLoginRequestOtp  (NEW)",
        [
            ["HIP method & path", "POST /v3/hip/verification/abha/profileLoginRequestOtp"],
            ["Downstream ABHA", "POST {AbhaNumberServiceUrl}/v3/profile/login/request/otp"],
            ["Controller", "VerificationController.ProfileLoginRequestOtp"],
            ["Purpose", "OTP for the chosen search-list index. Returns txnId to the UI (unlike most M1 OTP APIs)."],
        ],
    )
    add_para(doc, "HIP request body", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "scope": ["abha-login", "search-abha"],\n'
        '  "loginHint": "index",\n'
        '  "loginId": "1",\n'
        '  "otpSystem": "abdm",\n'
        '  "txnId": "<txnId from searchByMobile>"\n'
        "}",
    )
    add_para(
        doc,
        "Required: scope, loginHint, loginId, txnId. HIP encrypts loginId. Session TxnDictionary is updated "
        "with the new txnId from ABDM.",
        size=10,
    )
    add_para(doc, "HIP response (202 Accepted)", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "txnId": "uuid-from-abdm",\n'
        '  "message": "OTP sent to mobile ending xx89"\n'
        "}",
    )

    add_api_block(
        doc,
        "4.3 POST /v3/hip/verification/abha/profileLoginVerify  (NEW)",
        [
            ["HIP method & path", "POST /v3/hip/verification/abha/profileLoginVerify"],
            ["Downstream ABHA", "POST {AbhaNumberServiceUrl}/v3/profile/login/verify"],
            ["Extra ABDM header", "Transaction_Id = otp.txnId (passed through GatewayClient / HttpRequestHelper)"],
            ["Controller", "VerificationController.ProfileLoginVerify"],
            ["Purpose", "Verify OTP, store user X-token, return ABDM auth payload including token."],
        ],
    )
    add_para(doc, "HIP request body (plaintext OTP; HIP encrypts otpValue)", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "scope": ["abha-login", "search-abha"],\n'
        '  "authData": {\n'
        '    "authMethods": ["otp"],\n'
        '    "otp": {\n'
        '      "txnId": "<txnId from profileLoginRequestOtp>",\n'
        '      "otpValue": "123456"\n'
        "    }\n"
        "  }\n"
        "}",
    )
    add_para(doc, "HIP response (202 Accepted) when authResult is not failed", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "txnId": "...",\n'
        '  "authResult": "success",\n'
        '  "message": "OTP verified successfully",\n'
        '  "token": "<user-token>",\n'
        '  "expiresIn": 1800,\n'
        '  "refreshToken": "...",\n'
        '  "refreshExpiresIn": 1296000,\n'
        '  "accounts": [ { "abhaNumber": "91-xxxx-xxxx-xxxx", "preferredAbhaAddress": "...", "name": "..." } ]\n'
        "}",
    )
    add_para(
        doc,
        "On success HIP stores HealthIdNumberTokenDictionary[session_id] so getAbhaProfile and getAbhaCard work. "
        "Unlike /verification/verifyOtp, this API returns the ABDM token to the UI. If authResult is failed, "
        "HIP forwards ABDM status and body.",
        size=10,
        color=GRAY,
    )

    add_heading_styled(doc, "5. Behaviour change — existing requestOtp (MOBILE_NUMBER)", 1)
    add_para(
        doc,
        "Community mapped identifierType=MOBILE_NUMBER to loginHint=mobile and scopes [abha-login, mobile-verify]. "
        "Your branch maps it to loginHint=index and scopes [abha-login, mobile-verify, search-abha]. "
        "The identifier field is therefore treated as a list index, not a phone number.",
    )
    add_table(
        doc,
        ["Item", "Community abdm-v3", "Your upgraded branch"],
        [
            ["loginHint", "mobile", "index"],
            ["scope", "abha-login, mobile-verify", "abha-login, mobile-verify, search-abha"],
            ["identifier meaning", "mobile number (encrypted)", "ABHA list index (encrypted)"],
            ["Correct UI path for “find by mobile”", "This same requestOtp API", "New searchByMobile flow in §4"],
        ],
        col_widths=[4.0, 6.5, 6.5],
    )
    add_para(
        doc,
        "Frontend should not send a 10-digit mobile to POST /v3/hip/verification/requestOtp with MOBILE_NUMBER. "
        "Use the three new APIs. Keep requestOtp for ABHA_NUMBER and AADHAAR_NUMBER as before.",
        bold=True,
    )

    add_heading_styled(doc, "6. Fixes on existing create / login APIs", 1)
    add_heading_styled(doc, "6.1 Enrol stores txnId for address creation", 2)
    add_para(
        doc,
        "After POST /v3/hip/verifyOtpAndCreateABHA succeeds, HIP now AddOrUpdate TxnDictionary with "
        "enrollByAadhaarResponse.TxnId in addition to storing the user token. Without this, "
        "getAbhaAddressSuggestions / createAbhaAddress could still send the Aadhaar-OTP txnId and fail. "
        "Suggestions also refresh txnId from ABDM’s suggestion response.",
    )

    add_heading_styled(doc, "6.2 Profile fetch path (legacy helper)", 2)
    add_para(
        doc,
        "AbhaService.getABHAProfile now calls GET /v3/profile/account (ABHA_ACCOUNT) instead of "
        "GET /v2/account/profile (ABHA_PATIENT_PROFILE). Direct HIP route GET /v3/hip/verification/getAbhaProfile "
        "already used V3 in community; this aligns the shared helper used after login.",
    )

    add_heading_styled(doc, "6.3 Missing session token → 401", 2)
    add_para(
        doc,
        "GET /v3/hip/verification/getAbhaProfile now returns HTTP 401 with "
        "{ \"message\": \"Session token not found. Please complete OTP verification first.\" } "
        "when HealthIdNumberTokenDictionary has no token for the Bahmni session. Community could throw and 500.",
    )

    add_heading_styled(doc, "6.4 Login-verify JSON camelCase", 2)
    add_para(
        doc,
        "ABHALoginVerifyOTPRequest properties were PascalCase (AuthData, Scope, TxnId, OtpValue). "
        "ABDM V3 expects camelCase. Your branch serializes authData, scope, authMethods, otp, txnId, otpValue. "
        "This applies to POST /v3/hip/verification/verifyOtp (existing path), not only the new profileLoginVerify.",
    )
    add_para(doc, "HIP → ABHA body (verify OTP)", bold=True, space_after=4)
    add_code(
        doc,
        "{\n"
        '  "scope": ["abha-login", "aadhaar-verify"],\n'
        '  "authData": {\n'
        '    "authMethods": ["otp"],\n'
        '    "otp": {\n'
        '      "txnId": "<session txnId>",\n'
        '      "otpValue": "<RSA-encrypted OTP>"\n'
        "    }\n"
        "  }\n"
        "}",
    )

    add_heading_styled(doc, "6.5 Thread-safe session maps", 2)
    add_para(
        doc,
        "CreationMap dictionaries are ConcurrentDictionary. All M1 controllers use AddOrUpdate and TryGetValue. "
        "This avoids ArgumentException on duplicate keys and KeyNotFoundException under concurrent Bahmni users. "
        "Maps remain in-process memory (lost on HIP restart or another replica).",
    )
    add_table(
        doc,
        ["Dictionary", "Key", "Value", "M1 use"],
        [
            ["TxnDictionary", "Bahmni session_id", "Latest ABDM txnId", "OTP steps, search-by-mobile, address"],
            ["HealthIdNumberTokenDictionary", "session_id", "User X-token", "profile, card, verifyAbhaAccount"],
            ["HealthIdLoginScopeDictionary", "session_id", "Scope list", "verifyOtp / ABHA-address verifyOtp"],
        ],
        col_widths=[5.2, 3.5, 4.0, 4.3],
    )

    add_heading_styled(doc, "7. New files (M1)", 1)
    add_table(
        doc,
        ["File", "Role"],
        [
            ["Verification/Model/SearchAbhaByMobileRequest.cs", "scope + mobile"],
            ["Verification/Model/SearchAbhaByMobileResponse.cs", "txnId + abhaList (index, abhaNumber, name, gender)"],
            ["Verification/Model/ProfileLoginRequestOtpRequest.cs", "scope, loginHint, loginId, otpSystem, txnId"],
            ["Verification/Model/ProfileLoginVerifyRequest.cs", "scope + authData.otp"],
        ],
        col_widths=[8.5, 8.5],
    )

    add_heading_styled(doc, "8. UI contract — what to change vs community", 1)
    add_bullet(doc, "Add the three search-by-mobile screens/APIs for “patient knows only mobile”.")
    add_bullet(doc, "Keep using session-hidden txnId for create and ABHA-number login; search-by-mobile returns txnId on purpose.")
    add_bullet(doc, "After profileLoginVerify, existing getAbhaProfile / getAbhaCard work with the same Bahmni session.")
    add_bullet(doc, "Handle 401 on getAbhaProfile (complete OTP first).")
    add_bullet(doc, "Do not send raw mobile on /verification/requestOtp with MOBILE_NUMBER.")
    add_bullet(doc, "Always send CORRELATION-ID; HIP copies it to gateway and ABHA.")

    add_heading_styled(doc, "9. Code map (M1 deltas)", 1)
    add_table(
        doc,
        ["Area", "Primary types"],
        [
            ["New search/login APIs", "Verification/VerificationController.cs (SearchAbhaByMobile, ProfileLoginRequestOtp, ProfileLoginVerify)"],
            ["MOBILE_NUMBER mapping", "Verification/VerificationRequestMapper.cs"],
            ["Scopes / hints", "Common/Model/ABHAScope.cs, ABHALoginHint.cs"],
            ["Paths", "Common/Constants.cs (APP_PATH_ABHA_* and ABHA_SEARCH_BY_MOBILE)"],
            ["Session maps", "Creation/CreationMap.cs"],
            ["Create / card / txnId", "Creation/CreationController.cs, Creation/AbhaService.cs"],
            ["Verify JSON", "Verification/Model/ABHALoginVerifyOTPRequest.cs"],
        ],
        col_widths=[4.5, 12.5],
    )

    add_heading_styled(doc, "10. Related commits (upgraded repo, M1)", 1)
    add_bullet(doc, "83ddacfe — abdm V3 (baseline)")
    add_bullet(doc, "Multiple “M1 search by mobile” commits (0438c4f … 5f544a4)")
    add_bullet(doc, "4bd65bb / f88b89b — ABHA address creation fixes")
    add_bullet(doc, "136b63b — ABHA card download fixes")
    add_bullet(doc, "31723e3 / 448375d — ConcurrentDictionary / thread-safe maps")

    add_heading_styled(doc, "11. External references", 1)
    add_bullet(doc, "ABDM Sandbox V3: https://sandbox.abdm.gov.in/sandbox/v3/new-documentation")
    add_bullet(doc, "ABHA search: POST /v3/profile/account/abha/search")
    add_bullet(doc, "ABHA login OTP: POST /v3/profile/login/request/otp and /v3/profile/login/verify")
    add_bullet(doc, "User-initiated linking swagger (M2): sandbox v3 swagger integration_label abdm_user_initiated_linking_phr")
    add_bullet(doc, "HIP-initiated linking (M2): https://sandbox.abdm.gov.in/sandbox/v3/new-documentation?doc=HIPInitiatedlinking")

    footer = doc.add_paragraph()
    r = footer.add_run("End of M1 change notes. See V3ChangesExpalined-M2.docx for linking / HFR changes.")
    set_run(r, size=10, color=GRAY)
    footer.paragraph_format.space_before = Pt(18)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
